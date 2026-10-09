# Round 68 -- rebuilds the product-bearing READY files so every product code is one of the LIVE hyphenated codes ZV-001..ZV-010 (usage: build-r68-files.py <r65 fixtures dir> <out dir>).
# Live Product Master (screenshot): code, name, brand, strength, pack, SKU, division, UOM. ZV-002 was NOT visible and is INFERRED from the pattern. Rates are demo values.
import sys, shutil, openpyxl
from openpyxl.styles import PatternFill, Font
S, O = sys.argv[1], sys.argv[2]
# code, name, brand, strength, pack, sku, division, uom, group, category, PTR
P = [("ZV-001","Zivacard 10","Zivacard","10mg","1×5","SKU-ZV-001","Zivira","Tube","CV","CARDIO-VASCULAR",95),
     ("ZV-002","Zivabeta 5","Zivabeta","5mg","1×6","SKU-ZV-002","Astra","Strip","CV","CARDIO-VASCULAR",70),      # INFERRED (row not visible in the screenshot)
     ("ZV-003","Zivaflox 3","Zivaflox","3mg","1×7","SKU-ZV-003","Aura","Bottle","AI","ANTI-INFECTIVE",120),
     ("ZV-004","Zivalergy 2","Zivalergy","2mg","1×8","SKU-ZV-004","Zivira","Vial","AA","ANTI-ALLERGY",85),
     ("ZV-005","Zivaglauc 1","Zivaglauc","1mg","1×9","SKU-ZV-005","Astra","Box","AG","ANTI-GLAUCOMA",150),
     ("ZV-006","Zivapress 20","Zivapress","20mg","1×10","SKU-ZV-006","Aura","Tube","AH","ANTI-HYPERTENSIVE",110),
     ("ZV-007","Zivakid 15","Zivakid","15mg","1×11","SKU-ZV-007","Zivira","Strip","PD","PAEDIATRIC",60),
     ("ZV-008","Zivapain 8","Zivapain","8mg","1×12","SKU-ZV-008","Astra","Bottle","AN","ANALGESIC",45),
     ("ZV-009","Zivaskin 4","Zivaskin","4mg","1×13","SKU-ZV-009","Aura","Vial","DR","DERMATOLOGY",130),
     ("ZV-010","Zivaneuro 6","Zivaneuro","6mg","1×14","SKU-ZV-010","Zivira","Box","NE","NEURO",100)]
PTR = {p[0]: p[10] for p in P}
def hdr(path):
    ws = openpyxl.load_workbook(path).active
    return ws.title, [c.value for c in ws[1]], [c.value for c in ws[1] if c.fill and c.fill.fgColor and c.fill.fgColor.rgb in ("00FFFF00", "FFFFFF00")], [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r)]
def save(fn, sheet, headers, mandatory, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = sheet; ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        if c.value in mandatory: c.fill = PatternFill("solid", fgColor="FFFF00")
        ws.column_dimensions[c.column_letter].width = max(14, len(str(c.value)) + 4)
    for r in rows: ws.append(r)
    wb.save(f"{O}/{fn}"); print(fn, len(rows), "rows")
R = lambda n: f"{S}/{n}_READY_TO_UPLOAD.xlsx"
# 02 Product: template headers Product Code, Product Name, Group, Category, Brand, Pack, Division, Active (Strength / SKU / UOM are not template columns, so they stay as in the live master)
sh, h, m, _ = hdr(R("02_Product")); idx = {k: i for i, k in enumerate(h)}
rows = []
for p in P:
    r = [None] * len(h); r[idx["Product Code"]] = p[0]; r[idx["Product Name"]] = p[1]; r[idx["Group"]] = p[8]; r[idx["Category"]] = p[9]; r[idx["Brand"]] = p[2]; r[idx["Pack"]] = p[4]; r[idx["Division"]] = p[6]; r[idx["Active"]] = "Yes"; rows.append(r)
save("02_Product_READY_TO_UPLOAD.xlsx", sh, h, m, rows)
# 03 Product Rate
sh, h, m, _ = hdr(R("03_Product_Rate")); idx = {k: i for i, k in enumerate(h)}
rows = []
for p in P:
    r = [None] * len(h); r[idx["Product Code"]] = p[0]; r[idx["Product Name"]] = p[1]; r[idx["PTR"]] = p[10]; r[idx["PTS"]] = round(p[10] * .9, 2); r[idx["MRP"]] = round(p[10] * 1.25, 2); r[idx["Effective From"]] = "01/10/2026"; rows.append(r)
save("03_Product_Rate_READY_TO_UPLOAD.xlsx", sh, h, m, rows)
# 09 Target: same employees / months / quantities, product codes cycle through all ten live codes; rate = the product's PTR
ORDER = ["ZV-001", "ZV-006", "ZV-003", "ZV-004", "ZV-007", "ZV-005", "ZV-008", "ZV-009", "ZV-002", "ZV-010"]
sh, h, m, rows = hdr(R("09_Target")); idx = {k: i for i, k in enumerate(h)}
for i, r in enumerate(rows):
    c = ORDER[i % 10]; r[idx["Sale ERP Code"]] = c; r[idx["Target Rate"]] = PTR[c]; r[idx["Target Value"]] = PTR[c] * r[idx["Target Qty"]]
save("09_Target_READY_TO_UPLOAD.xlsx", sh, h, m, rows)
# 10 Sample Despatch
MAP = {"ZV024": "ZV-001", "ZV026": "ZV-006", "ZV013": "ZV-003", "ZV010": "ZV-004"}
sh, h, m, rows = hdr(R("10_Sample_Despatch")); i = h.index("Sample ERP Code")
for r in rows: r[i] = MAP[r[i]]
save("10_Sample_Despatch_READY_TO_UPLOAD.xlsx", sh, h, m, rows)
# the user's own despatch template (6 sample rows)
save("Upl_Despatch_Master_SAMPLE.xlsx", "Upl_Despatch_Master", ["Employee ID", "Sample ERP Code", "Despatch Qty"], [],
     [["MR-001", "ZV-001", 100], ["MR-001", "ZV-006", 75], ["MR-002", "ZV-003", 50], ["MR-002", "ZV-004", 100], ["MR-003", "ZV-001", 25], ["MR-003", "ZV-006", 60]])
for n in ["01_Salesforce_OPTIONAL", "04_Stockist", "05_Listed_Doctor", "06_Chemists", "07_Holiday", "08_Leave", "11_Input_Despatch"]:
    shutil.copy(R(n), f"{O}/{n}_READY_TO_UPLOAD.xlsx"); print(n, "unchanged")
